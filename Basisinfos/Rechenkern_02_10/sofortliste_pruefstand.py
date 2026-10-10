"""Pruefstand Sofortliste S-1 bis S-5 (GUI-Bereinigung, Plan Schritt 53 G-E, 10.10.2026) - Views EINZELN gegen eine
Wegwerf-DB (nie main.py am Desktop, Memory feedback_desktop_kein_produktivstart), kein Netz, keine Mail.

    python Basisinfos/Rechenkern_02_10/sofortliste_pruefstand.py

  S1a compute_cost_basis_view ohne `menge` rechnet wie bisher (Analysten unveraendert)
  S1b mit `menge` = frei + gestakt: voll gestakt bekommt Wert und G/V
  S1c PortfolioView-Zeile: Wert = (frei + gestakt) x Preis; Gesamtwert unveraendert (Summen waren schon richtig); Zusatz ohne
      "im Regelwerk noch nicht beruecksichtigt"
  S2  Einstandspreise berechnen: beim NEUEN Abgleich bleibt staked_quantity unberuehrt, beim ALTEN wird geschrieben (Seiteneffekt)
  S3a Export: Spalten Gestakt und Gesamt; Rundweg - der Import liest "Anzahl Coins" = frei (keine Doppelzaehlung)
  S3b Import: ein vom Bitpanda-Abgleich gefuehrter Bestand wird NICHT ueberschrieben (Warnung), ein Import-Bestand schon
  S4  SignalsView zeigt den Hinweis NICHT AKTUELL, solange die Kette angehalten ist - und nicht, wenn sie laeuft
  S5  Sperrtext nennt den Halt; ohne Halt der alte Text; die alten irrefuehrenden Saetze sind aus den Dateien verschwunden
  S6  Standard-DB data/tradinginfotool.db unveraendert (Pruefsumme)
"""
import hashlib
import io
import os
import sqlite3
import sys
import tempfile
from datetime import datetime, timezone

WURZEL = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, WURZEL)
os.chdir(WURZEL)
STD = os.path.join("data", "tradinginfotool.db")
vorher = hashlib.sha256(open(STD, "rb").read()).hexdigest() if os.path.exists(STD) else None

import config  # noqa: E402
import database.db as DB  # noqa: E402
from database.models import Holding, PriceSnapshot  # noqa: E402
import importer.bitpanda_avg_cost as AC  # noqa: E402
import importer.excel_import as XI  # noqa: E402
from scheduler import rollen_job as RJ  # noqa: E402
from agent import regel0_groesse as G  # noqa: E402

ok = n = 0


def pruefe(name, gut, info=""):
    global ok, n
    n += 1; ok += bool(gut)
    print("%-4s %s  %s" % (name, "ok" if gut else "FEHLER", info), flush=True)


def holding(sym, q, st, avg=None, tracked=None, manual=None):
    h = Holding.__new__(Holding)
    for k, v in dict(symbol=sym, quantity=q, staked_quantity=st, avg_buy_price_eur=avg, avg_buy_price_tracked_qty=tracked,
                     avg_buy_price_manual_eur=manual, updated_at="2026-10-10T00:00:00+00:00", source="bitpanda_sync",
                     avg_buy_price_computed_at=None).items():
        object.__setattr__(h, k, v) if hasattr(type(h), "__dataclass_fields__") else setattr(h, k, v)
    return h


# ---- S1a / S1b
try:
    h = DB.Holding if hasattr(DB, "Holding") else Holding
    hx = Holding(symbol="ETH", quantity=10.0, updated_at="x", source="bitpanda_sync", staked_quantity=5.0,
                 avg_buy_price_eur=2.0, avg_buy_price_tracked_qty=12.0)
except TypeError:
    hx = holding("ETH", 10.0, 5.0, 2.0, 12.0)
alt = AC.compute_cost_basis_view(hx, 3.0)
neu = AC.compute_cost_basis_view(hx, 3.0, menge=15.0)
pruefe("S1a", alt.known_quantity == 10.0 and alt.current_value_eur == 30.0 and abs(alt.pl_pct - 50.0) < 1e-9,
       "ohne menge: bekannt %s, Wert %s, G/V %s" % (alt.known_quantity, alt.current_value_eur, alt.pl_pct))
try:
    hv = Holding(symbol="ALGO", quantity=0.0, updated_at="x", source="bitpanda_sync", staked_quantity=100.0,
                 avg_buy_price_eur=0.2, avg_buy_price_tracked_qty=100.0)
except TypeError:
    hv = holding("ALGO", 0.0, 100.0, 0.2, 100.0)
v0 = AC.compute_cost_basis_view(hv, 0.1)
v1 = AC.compute_cost_basis_view(hv, 0.1, menge=100.0)
pruefe("S1b", neu.known_quantity == 12.0 and neu.unknown_quantity == 3.0 and neu.current_value_eur == 45.0
       and v0.current_value_eur == 0.0 and v0.pl_pct is None and abs(v1.current_value_eur - 10.0) < 1e-9 and abs(v1.pl_pct + 50.0) < 1e-9,
       "mit menge: ETH bekannt %s/unbekannt %s, Wert %s · ALGO voll gestakt Wert %s -> %s, G/V %s -> %s"
       % (neu.known_quantity, neu.unknown_quantity, neu.current_value_eur, v0.current_value_eur, v1.current_value_eur, v0.pl_pct, v1.pl_pct))

# ---- Wegwerf-DB
tmp = tempfile.mkdtemp(prefix="sofortliste_")
pfad = os.path.join(tmp, "wegwerf.db")
c0 = sqlite3.connect(pfad)
c0.row_factory = sqlite3.Row
DB.init_db(c0)
for sym, q, st in (("ETH", 0.05685134, 0.48714122), ("ALGO", 0.0, 2694.66), ("BTC", 0.05496596, 0.0)):
    DB.upsert_holding(c0, sym, q, source="bitpanda_sync")
    DB.update_holding_staked_quantity(c0, sym, st)
jetzt = datetime.now(timezone.utc).isoformat()
for sym, cg, p in (("ETH", "ethereum", 2226.23), ("ALGO", "algorand", 0.103278), ("BTC", "bitcoin", 73831.0)):
    DB.insert_price_snapshot(c0, PriceSnapshot(symbol=sym, coingecko_id=cg, price_usd=p * 1.12, price_eur=p, market_cap_usd=None,
                                               volume_24h_usd=None, change_24h_pct=None, fetched_at=jetzt))
c0.commit()


def fabrik():
    c = sqlite3.connect(pfad)
    c.row_factory = sqlite3.Row
    return c


# ---- S1c PortfolioView
import tkinter as tk  # noqa: E402
root = tk.Tk()
root.withdraw()
from ui.portfolio import PortfolioView  # noqa: E402
wl = config.get_watchlist()
pv = PortfolioView(root, fabrik, wl)
pv.refresh() if hasattr(pv, "refresh") else None
zeilen = {iid: pv.tree.item(iid, "values") for iid in pv.tree.get_children()}
algo = zeilen.get("ALGO")
eth = zeilen.get("ETH")
gesamt = pv.total_label.cget("text")
soll_algo = 2694.66 * 0.103278
soll_eth = (0.05685134 + 0.48714122) * 2226.23


def zahl(t):
    return float(str(t).replace(",", "").replace(" EUR", "").strip())


soll_gesamt = 0.05685134 * 2226.23 + 0.05496596 * 73831.0 + 0.48714122 * 2226.23 + 2694.66 * 0.103278
gl = float(gesamt.split("Gesamtwert:")[1].split("EUR")[0].replace(",", "").strip())
pruefe("S1c", algo is not None and abs(zahl(algo[5]) - soll_algo) < 0.01 and abs(zahl(eth[5]) - soll_eth) < 0.01
       and abs(gl - soll_gesamt - (pv._cash_fiat if hasattr(pv, "_cash_fiat") else 0)) < 1.0 and "im Regelwerk" not in gesamt,
       "ALGO Zeile %s (Soll %.2f) · ETH %s (Soll %.2f) · %s" % (algo[5] if algo else "-", soll_algo, eth[5] if eth else "-", soll_eth, gesamt))

# ---- S2 Einstandspreise berechnen
alt_get, alt_comp, alt_quelle = AC.get_wallet_transactions, AC.compute_staked_quantities, config.bitpanda_bestand_quelle
try:
    AC.get_wallet_transactions = lambda *a, **k: []
    AC.compute_staked_quantities = lambda tx, existing=None: {"ETH": 0.9711}      # die alte Rekonstruktion (verdoppelt)
    erg = {}
    for quelle in ("neu", "alt"):
        config.bitpanda_bestand_quelle = (lambda q=quelle: q)
        c = fabrik()
        DB.update_holding_staked_quantity(c, "ETH", 0.48714122)
        c.commit()
        AC.sync_avg_buy_prices(c, "k", wl, [])
        erg[quelle] = c.execute("SELECT staked_quantity FROM holdings WHERE symbol='ETH'").fetchone()[0]
        c.close()
finally:
    AC.get_wallet_transactions, AC.compute_staked_quantities, config.bitpanda_bestand_quelle = alt_get, alt_comp, alt_quelle
pruefe("S2", abs(erg["neu"] - 0.48714122) < 1e-12 and abs(erg["alt"] - 0.9711) < 1e-12,
       "staked ETH nach dem Menue: neuer Abgleich %s (unberuehrt) · alter Abgleich %s (geschrieben)" % (erg["neu"], erg["alt"]))

# S2 im ALTEN Weg setzt - wie vorgesehen - gestakte Mengen zurueck, die die alte Rekonstruktion nicht kennt (ALGO -> 0).
# Fuer S3 den Ausgangsstand der Wegwerf-DB wiederherstellen.
c = fabrik()
for sym, st in (("ETH", 0.48714122), ("ALGO", 2694.66), ("BTC", 0.0)):
    DB.update_holding_staked_quantity(c, sym, st)
c.commit(); c.close()

# ---- S3a Export / Rundweg
from pathlib import Path  # noqa: E402
xp = Path(os.path.join(tmp, "export.xlsx"))
c = fabrik()
XI.export_holdings(c, xp)
c.close()
import openpyxl  # noqa: E402
wb = openpyxl.load_workbook(xp, data_only=True)
kopf = [x.value for x in next(wb[XI.SHEET_NAME_KRYPTO].iter_rows(min_row=1, max_row=1))]
reihe = {r[1].value: [x.value for x in r] for r in wb[XI.SHEET_NAME_KRYPTO].iter_rows(min_row=2) if r[1].value}
rund = XI.read_holdings_from_excel(xp)
pruefe("S3a", kopf[:6] == ["Coin Name", "Kurzzeichen", "Anzahl Coins", "Quelle", "Gestakt", "Gesamt"]
       and reihe.get("ALGO", [None] * 6)[4] == 2694.66 and reihe["ALGO"][2] == 0.0 and abs(reihe["ETH"][5] - 0.54399256) < 1e-9
       and rund.get("ALGO") == 0.0 and abs(rund.get("ETH") - 0.05685134) < 1e-12,
       "Kopf %s · ALGO frei %s gestakt %s · Rundweg liest frei: ALGO %s, ETH %s" % (kopf, reihe.get("ALGO", ["-"] * 6)[2],
                                                                                  reihe.get("ALGO", ["-"] * 6)[4], rund.get("ALGO"), rund.get("ETH")))

# ---- S3b Import schuetzt Abgleich-Bestaende
ws = wb[XI.SHEET_NAME_KRYPTO]
for r in ws.iter_rows(min_row=2):
    if r[1].value == "BTC":
        r[2].value = 9.0
wb.save(xp)
c = fabrik()
c.execute("UPDATE holdings SET source='import' WHERE symbol='ETH'")
c.commit()
for r in openpyxl.load_workbook(xp)[XI.SHEET_NAME_KRYPTO].iter_rows(min_row=2):
    pass
wb2 = openpyxl.load_workbook(xp)
for r in wb2[XI.SHEET_NAME_KRYPTO].iter_rows(min_row=2):
    if r[1].value == "ETH":
        r[2].value = 7.0
wb2.save(xp)
res = XI.import_holdings(c, xp)
btc = c.execute("SELECT quantity FROM holdings WHERE symbol='BTC'").fetchone()[0]
ethq = c.execute("SELECT quantity FROM holdings WHERE symbol='ETH'").fetchone()[0]
c.close()
pruefe("S3b", abs(btc - 0.05496596) < 1e-12 and ethq == 7.0 and any("BTC" in w and "NICHT übernommen" in w for w in res.warnings),
       "BTC (Abgleich) bleibt %s, ETH (Import) wird %s · Warnung: %s" % (btc, ethq, [w for w in res.warnings if "BTC" in w][:1]))

# ---- S4 SignalsView
from ui.signals_view import SignalsView  # noqa: E402


def banner(view):
    return [w.cget("text") for w in view.winfo_children() if w.winfo_class() == "TLabel" and "NICHT AKTUELL" in str(w.cget("text"))]


sv = SignalsView(root, fabrik, wl, None, None, None)
b_an = banner(sv)
alt_lade = G.lade
try:
    G.lade = lambda pfad=None: dict(alt_lade(pfad), spot_kette_angehalten=False)
    sv2 = SignalsView(root, fabrik, wl, None, None, None)
    b_aus = banner(sv2)
    txt_aus = RJ.alte_analyse_hinweis("krypto")
finally:
    G.lade = alt_lade
pruefe("S4", len(b_an) == 1 and "05.10.2026" in b_an[0] and not b_aus, "Hinweis bei Halt: %s · ohne Halt: %d" % (b_an[:1], len(b_aus)))

# ---- S5 Texte
txt_an = RJ.alte_analyse_hinweis("krypto")
quellen = {p: io.open(p, encoding="utf-8").read() for p in ("ui/app.py", "ui/regime_view.py", "ui/hebel_view.py", "ui/screener_view.py",
                                                         "ui/portfolio.py")}
weg = ["Die Führung offener Positionen kommt später (O13)", "Wird automatisch im nächsten Budget-Allocator-Lauf",
       "Wirkt ab dem nächsten Pipeline-Lauf", "normale Signal-Pipeline, nach einem App-Neustart", "(EUR, manuell)",
       "im Regelwerk noch nicht berücksichtigt", "An = die Rollen-Kette darf für dieses Asset ein Hebelgeschäft rechnen",
       "Wirkt nach Commit+Push (Desktop) "]
rest = [w for w in weg if any(w in q for q in quellen.values())]
pruefe("S5", "angehalten" in (txt_an or "") and "automatisch im" not in (txt_an or "") and "automatisch im" in (txt_aus or "") and not rest,
       "Sperrtext bei Halt: %r · alte Saetze noch da: %s" % ((txt_an or "")[:90], rest))
root.destroy()

# ---- S6 Standard-DB
nachher = hashlib.sha256(open(STD, "rb").read()).hexdigest() if os.path.exists(STD) else None
pruefe("S6", vorher == nachher, "data/tradinginfotool.db %s" % ("unveraendert" if vorher == nachher else "VERAENDERT"))
print("\n%d von %d bestanden" % (ok, n))
